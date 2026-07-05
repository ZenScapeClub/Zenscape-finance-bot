import os, json, csv, io, base64, logging, tempfile
from datetime import datetime, timedelta
from enum import Enum

from telegram import Update, InlineKeyboardButton, InlineKeyboardMarkup
from telegram.ext import (Application, CommandHandler, MessageHandler,
    CallbackQueryHandler, ConversationHandler, ContextTypes, filters)
from telegram.constants import ParseMode
import gspread
from google.oauth2.service_account import Credentials

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

# ── Access control ────────────────────────────────────────────────────────────

def _parse_ids(env_var, fallback=''):
    val = os.environ.get(env_var, fallback)
    return {int(x.strip()) for x in val.split(',') if x.strip().isdigit()}

ADMIN_USERS  = _parse_ids('ALLOWED_USERS', '257170336')
VIEWER_USERS = _parse_ids('VIEWER_USERS')

def admin_only(func):
    async def wrapper(update: Update, context: ContextTypes.DEFAULT_TYPE):
        uid = update.effective_user.id
        if uid not in ADMIN_USERS:
            await update.effective_message.reply_text("Нет доступа.")
            return
        return await func(update, context)
    wrapper.__name__ = func.__name__
    return wrapper

def any_user(func):
    async def wrapper(update: Update, context: ContextTypes.DEFAULT_TYPE):
        uid = update.effective_user.id
        if uid not in ADMIN_USERS and uid not in VIEWER_USERS:
            await update.effective_message.reply_text("Нет доступа.")
            return
        return await func(update, context)
    wrapper.__name__ = func.__name__
    return wrapper

def _is_admin(update: Update) -> bool:
    return update.effective_user.id in ADMIN_USERS

# ── States ────────────────────────────────────────────────────────────────────

class S(Enum):
    MENU=1; SELECT_PROJECT=2; OP_TYPE=3; CATEGORY=4; AMOUNT=5
    DATE=6; PAY_STATUS=7; CONTRACTOR=8; COMMENT=9
    CP_NAME=20; CP_REVENUE=21; CP_EXPENSE=22
    EDIT_LIST=30
    IMPORT_FILE=40; IMPORT_CONFIRM=41
    QUICK_PROJECT=50; QUICK_AMOUNT=51

# ── Categories ────────────────────────────────────────────────────────────────

EXPENSE_CATS = ['Озеленение', 'Оплата подряду', 'Логистика', 'Строительные',
                'Командировка', 'Химия и уход', 'ГСМ', 'Расходы на ЗСД', 'Прочее']
INCOME_CATS  = ['Поступление от клиента', 'Возврат', 'Прочие доходы']

CAT_MAP = {
    'Логитстика': 'Логистика',
    'Оплата субподряду': 'Оплата подряду',
    'Оплата подрядчикам': 'Оплата подряду',
    'Оплата субподряд': 'Оплата подряду',
    'Строительство': 'Строительные',
    'Поступление от клиентов': 'Поступление от клиента',
    'Поступление от клиента': 'Поступление от клиента',
    'Растения': 'Озеленение',
    'Агентское': 'Прочее',
}

def _norm_cat(cat):
    if not cat:
        return 'Прочее'
    cat = cat.strip()
    return CAT_MAP.get(cat, cat)

# ── Sheets ────────────────────────────────────────────────────────────────────

_sheets = None

def get_sheets():
    global _sheets
    if _sheets is None:
        b64 = os.environ['GOOGLE_SERVICE_ACCOUNT_B64']
        sa = json.loads(base64.b64decode(b64).decode())
        creds = Credentials.from_service_account_info(
            sa, scopes=['https://www.googleapis.com/auth/spreadsheets'])
        _sheets = gspread.authorize(creds).open_by_key(os.environ['GOOGLE_SHEETS_ID'])
        _ensure_base_sheets(_sheets)
    return _sheets

def _ws_titles(sp):
    return [ws.title for ws in sp.worksheets()]

def _ensure_base_sheets(sp):
    titles = _ws_titles(sp)
    if 'Объекты' not in titles:
        ws = sp.add_worksheet('Объекты', 100, 10)
        ws.update('A1:J1', [['Название','Статус','Дата начала','Адрес',
            'Сумма договора','План расход','План прибыль','План маржа %',
            'Факт доход','Факт расход']])
    if 'Операции' not in titles:
        ws = sp.add_worksheet('Операции', 1000, 9)
        ws.update('A1:I1', [['ID','Дата','Объект','Тип','Категория',
            'Сумма','Контрагент','Статус оплаты','Комментарий']])
    if 'Дашборд' not in titles:
        sp.add_worksheet('Дашборд', 200, 12)
    if 'Cash Flow' not in titles:
        sp.add_worksheet('Cash Flow', 20, 30)

# ── Supabase mirror ───────────────────────────────────────────────────────────
# Дублируем записи в облачную БД ZenScape (единый источник для приложения).
# Если переменные окружения не заданы — тихо пропускаем, бот работает как раньше.

import requests as _rq

SB_URL = os.environ.get('SUPABASE_URL', '').rstrip('/')
SB_KEY = os.environ.get('SUPABASE_SERVICE_KEY', '')

# Категории бота → статьи приложения ZenScape
SB_ARTICLE_MAP = {
    'Поступление от клиента': ('income.client',      'Доходы'),
    'Возврат':                ('income.refund',      'Возврат доходы'),
    'Прочие доходы':          ('income.other',       'Доходы'),
    'Озеленение':             ('direct.planting',    'Прямые расходы'),
    'Строительные':           ('direct.construction','Прямые расходы'),
    'Оплата подряду':         ('direct.cost',        'Прямые расходы'),
    'Логистика':              ('direct.project',     'Прямые расходы'),
    'Командировка':           ('travel.fare',        'Командировки'),
    'Химия и уход':           ('direct.maintenance', 'Прямые расходы'),
    'ГСМ':                    ('transport.fuel',     'Транспорт'),
    'Расходы на ЗСД':         ('transport.zsd',      'Транспорт'),
    'Прочее':                 ('overhead.other',     'Общехозяйственные'),
}

def _sb_enabled():
    return bool(SB_URL and SB_KEY)

def _sb_headers():
    return {'apikey': SB_KEY, 'Authorization': f'Bearer {SB_KEY}',
            'Content-Type': 'application/json', 'Prefer': 'return=representation'}

def _sb_project_id(name):
    try:
        r = _rq.get(f'{SB_URL}/rest/v1/projects', headers=_sb_headers(),
                    params={'name': f'eq.{name}', 'select': 'id'}, timeout=10)
        rows = r.json()
        return rows[0]['id'] if rows else None
    except Exception as e:
        logger.warning(f'Supabase project lookup failed: {e}')
        return None

def _sb_date(d):
    try:
        return datetime.strptime(d, '%d.%m.%Y').strftime('%Y-%m-%d')
    except Exception:
        return d

def sb_mirror_operation(op_id, project, op_type, category, amount, date, pay_status, contractor='', comment=''):
    if not _sb_enabled():
        return
    try:
        pid = _sb_project_id(project)
        if not pid:
            logger.warning(f'Supabase: объект «{project}» не найден, операция {op_id} не зеркалирована')
            return
        cat = _norm_cat(category)
        article, group = SB_ARTICLE_MAP.get(cat, ('overhead.other', 'Общехозяйственные'))
        if op_type == 'Приход' and not article.startswith('income.'):
            article, group = 'income.client', 'Доходы'
        s = abs(float(amount))
        if op_type == 'Расход':
            s = -s
        purpose = comment or ''
        if article in ('direct.cost', 'direct.project', 'overhead.other') and cat not in ('Прочее',):
            purpose = f'[{cat}] {purpose}'.strip()
        _rq.post(f'{SB_URL}/rest/v1/expenses', headers=_sb_headers(), timeout=10, json={
            'project_id': pid, 'date': _sb_date(date), 'sum': s,
            'article': article, 'article_group': group,
            'purpose': purpose, 'counterparty': contractor or '',
            'plan_fact': 'plan' if pay_status == 'Ожидает' else 'fact',
            'sort_order': int(op_id),
        })
    except Exception as e:
        logger.warning(f'Supabase mirror add failed: {e}')

def sb_mirror_project(name, revenue):
    if not _sb_enabled():
        return
    try:
        r = _rq.get(f'{SB_URL}/rest/v1/studios', headers=_sb_headers(),
                    params={'select': 'id', 'limit': '1'}, timeout=10)
        studios = r.json()
        if not studios:
            return
        _rq.post(f'{SB_URL}/rest/v1/projects', headers=_sb_headers(), timeout=10, json={
            'studio_id': studios[0]['id'], 'name': name, 'status': 'in_progress',
            'total_budget': revenue, 'started_at': datetime.now().strftime('%Y-%m-%d'),
        })
    except Exception as e:
        logger.warning(f'Supabase mirror project failed: {e}')

def sb_mirror_delete(op_id):
    if not _sb_enabled():
        return
    try:
        _rq.delete(f'{SB_URL}/rest/v1/expenses', headers=_sb_headers(),
                   params={'sort_order': f'eq.{op_id}'}, timeout=10)
    except Exception as e:
        logger.warning(f'Supabase mirror delete failed: {e}')

# ── Обратная синхронизация: облако → Google Sheets ────────────────────────────
# Раз в 3 минуты подтягиваем то, что появилось в облаке с сайта:
#   • новые объекты (проекты) → лист «Объекты»
#   • новые операции (sort_order = 0) → лист «Операции» с присвоением ID

# Статьи приложения → категории бота
SB_ARTICLE_REVERSE = {
    'income.client':      ('Приход', 'Поступление от клиента'),
    'income.refund':      ('Приход', 'Возврат'),
    'income.other':       ('Приход', 'Прочие доходы'),
    'income.agent':       ('Приход', 'Прочие доходы'),
    'income.purchases':   ('Приход', 'Прочие доходы'),
    'income.supervision': ('Приход', 'Прочие доходы'),
    'income.adjust':      ('Приход', 'Прочие доходы'),
    'direct.planting':    ('Расход', 'Озеленение'),
    'direct.construction':('Расход', 'Строительные'),
    'direct.cost':        ('Расход', 'Оплата подряду'),
    'direct.project':     ('Расход', 'Логистика'),
    'direct.maintenance': ('Расход', 'Химия и уход'),
    'travel.fare':        ('Расход', 'Командировка'),
    'transport.fuel':     ('Расход', 'ГСМ'),
    'transport.zsd':      ('Расход', 'Расходы на ЗСД'),
}

def _sheet_date(iso):
    try:
        return datetime.strptime(str(iso)[:10], '%Y-%m-%d').strftime('%d.%m.%Y')
    except Exception:
        return str(iso)

def sync_from_supabase():
    """Однократный проход синхронизации облако → Sheets."""
    if not _sb_enabled():
        return
    sp = get_sheets()

    # 1. Новые объекты
    r = _rq.get(f'{SB_URL}/rest/v1/projects', headers=_sb_headers(),
                params={'select': 'id,name,status,total_budget,started_at'}, timeout=15)
    cloud_projects = r.json() if r.ok else []
    ws_obj = sp.worksheet('Объекты')
    sheet_names = {row[0].strip() for row in ws_obj.get_all_values()[1:] if row and row[0]}
    added_projects = []
    for cp in cloud_projects:
        name = (cp.get('name') or '').strip()
        if not name or name in sheet_names:
            continue
        budget = float(cp.get('total_budget') or 0)
        ws_obj.append_row([name, 'Активный',
                           _sheet_date(cp.get('started_at') or datetime.now().strftime('%Y-%m-%d')),
                           '', budget, 0, budget, 100.0 if budget else 0, 0, 0])
        _create_object_sheet(sp, name)
        added_projects.append(name)
        logger.info(f'Sync: новый объект с сайта → Sheets: {name}')

    # 2. Новые операции с сайта (sort_order = 0)
    r = _rq.get(f'{SB_URL}/rest/v1/expenses', headers=_sb_headers(),
                params={'select': 'id,project_id,date,sum,article,purpose,counterparty,plan_fact',
                        'sort_order': 'eq.0', 'order': 'created_at.asc'}, timeout=15)
    site_ops = r.json() if r.ok else []
    if not site_ops and not added_projects:
        return

    proj_name = {p['id']: p['name'] for p in cloud_projects}
    ws_ops = sp.worksheet('Операции')
    touched = set(added_projects)
    for op in site_ops:
        pname = proj_name.get(op.get('project_id'))
        if not pname:
            continue
        s = float(op.get('sum') or 0)
        op_type, category = SB_ARTICLE_REVERSE.get(
            op.get('article') or '',
            ('Приход', 'Прочие доходы') if s > 0 else ('Расход', 'Прочее'))
        next_id = _next_op_id(sp)
        ws_ops.append_row([next_id, _sheet_date(op.get('date')), pname, op_type,
                           category, abs(s), op.get('counterparty') or '',
                           'Ожидает' if op.get('plan_fact') == 'plan' else 'Оплачено',
                           op.get('purpose') or ''])
        _rq.patch(f'{SB_URL}/rest/v1/expenses', headers=_sb_headers(),
                  params={'id': f"eq.{op['id']}"},
                  json={'sort_order': next_id}, timeout=10)
        touched.add(pname)
        logger.info(f'Sync: операция с сайта → Sheets: #{next_id} {pname} {s}')

    for pname in touched:
        try:
            _update_object_sheet(sp, pname)
        except Exception as e:
            logger.warning(f'Sync: object sheet {pname}: {e}')
    if touched:
        _refresh_dashboard(sp)
        _refresh_cashflow(sp)

def _sync_loop():
    import time as _time
    _time.sleep(30)  # даём боту стартовать
    while True:
        try:
            sync_from_supabase()
        except Exception as e:
            logger.warning(f'Sync loop error: {e}')
        _time.sleep(180)

def start_sync_thread():
    if not _sb_enabled():
        logger.info('Sync: SUPABASE_URL/KEY не заданы — фоновая синхронизация выключена')
        return
    import threading
    threading.Thread(target=_sync_loop, daemon=True, name='sb-sync').start()
    logger.info('Sync: фоновая синхронизация облако → Sheets запущена (каждые 3 мин)')

# ── Raw data helpers ──────────────────────────────────────────────────────────

def _num(val):
    try:
        return float(str(val).replace(' ','').replace(',','.').replace('\xa0','').replace('₽','').strip())
    except:
        return 0.0

def _fmt(v):
    try: return f"{float(v):,.0f}".replace(',', ' ')
    except: return str(v) if v else "0"

def _pct(v):
    try: return f"{float(v)*100:.1f}%"
    except: return "—"

def _parse_date(s):
    for fmt in ('%d.%m.%Y', '%Y-%m-%d'):
        try:
            return datetime.strptime(s, fmt)
        except:
            pass
    return None

def _get_projects_raw(sp):
    try:
        rows = sp.worksheet('Объекты').get_all_values()[1:]
        return [r[0].strip() for r in rows if r and r[0].strip()]
    except:
        return []

def _get_project_row_raw(sp, project):
    try:
        rows = sp.worksheet('Объекты').get_all_values()[1:]
        for i, r in enumerate(rows):
            if r and r[0].strip() == project:
                return i + 2, r
        return None, None
    except:
        return None, None

def _get_operations_raw(sp):
    try:
        rows = sp.worksheet('Операции').get_all_values()[1:]
        ops = []
        for r in rows:
            if r and r[0]:
                ops.append({'id':r[0],'date':r[1],'project':r[2],'type':r[3],
                    'category':r[4],'amount':_num(r[5]),
                    'contractor':r[6] if len(r)>6 else '',
                    'pay_status':r[7] if len(r)>7 else '',
                    'comment':r[8] if len(r)>8 else ''})
        return ops
    except:
        return []

def _next_op_id(sp):
    rows = sp.worksheet('Операции').get_all_values()[1:]
    ids = [int(r[0]) for r in rows if r and r[0] and str(r[0]).isdigit()]
    return max(ids) + 1 if ids else 1

# ── Object sheet ──────────────────────────────────────────────────────────────

def _create_object_sheet(sp, name):
    try:
        ws = sp.add_worksheet(name[:50], 200, 10)
        _update_object_sheet(sp, name, ws)
    except Exception as e:
        logger.error(f"Error creating sheet for {name}: {e}")

def _update_object_sheet(sp, name, ws=None):
    try:
        if ws is None:
            try:
                ws = sp.worksheet(name[:50])
            except:
                ws = sp.add_worksheet(name[:50], 200, 10)

        ops = _get_operations_raw(sp)
        p_ops = [o for o in ops if o['project'] == name]
        _, row = _get_project_row_raw(sp, name)

        plan_rev = _num(row[4]) if row and len(row) > 4 else 0
        plan_exp = _num(row[5]) if row and len(row) > 5 else 0
        plan_profit = plan_rev - plan_exp
        plan_margin = plan_profit / plan_rev if plan_rev > 0 else 0

        f_inc = sum(o['amount'] for o in p_ops if o['type'] == 'Приход')
        f_exp = sum(o['amount'] for o in p_ops if o['type'] == 'Расход')
        f_profit = f_inc - f_exp
        f_margin = f_profit / f_inc if f_inc > 0 else 0

        ws.clear()
        ws.update('A1:B1', [['Объект:', name]])
        ws.update('A2:B7', [
            ['Статус', row[1] if row and len(row)>1 else ''],
            ['Сумма договора', plan_rev],
            ['План расход', plan_exp],
            ['План прибыль', plan_profit],
            ['План маржа %', round(plan_margin * 100, 1)],
            ['Маржа (мой доход)', f_profit],
        ])
        ws.update('D2:E7', [
            ['Факт доход', f_inc],
            ['Факт расход', f_exp],
            ['Факт прибыль', f_profit],
            ['Факт маржа %', round(f_margin * 100, 1)],
            ['Откл. прибыль', f_profit - plan_profit],
            ['', ''],
        ])
        ws.update('A9:I9', [['ID','Дата','Тип','Категория','Сумма',
                              'Контрагент','Статус оплаты','Комментарий','']])
        if p_ops:
            rows_data = [[o['id'], o['date'], o['type'], o['category'],
                          o['amount'], o['contractor'], o['pay_status'], o['comment'], '']
                         for o in sorted(p_ops, key=lambda x: x['date'], reverse=True)]
            ws.update(f'A10:I{9+len(rows_data)}', rows_data)
    except Exception as e:
        logger.error(f"Error updating sheet {name}: {e}")

# ── Dashboard ─────────────────────────────────────────────────────────────────

def _refresh_dashboard(sp):
    try:
        ws = sp.worksheet('Дашборд')
        ws.clear()
        projects = _get_projects_raw(sp)
        ops = _get_operations_raw(sp)

        all_income  = sum(o['amount'] for o in ops if o['type'] == 'Приход')
        all_expense = sum(o['amount'] for o in ops if o['type'] == 'Расход')
        all_profit  = all_income - all_expense
        all_margin  = all_profit / all_income if all_income > 0 else 0

        ws.update('A1', [['ZenScape — Финансовый дашборд']])
        ws.update('A2:B7', [
            ['Всего объектов', len(projects)],
            ['Факт доход', all_income],
            ['Факт расход', all_expense],
            ['Факт прибыль', all_profit],
            ['Факт маржа %', round(all_margin * 100, 1)],
            ['Маржа (мой доход)', all_profit],
        ])

        ws.update('A9:L9', [['Объект','Статус','Договор','План расход',
            'План прибыль','План маржа %','Факт доход','Факт расход',
            'Факт прибыль','Факт маржа %','Откл. прибыль','Маржа ₽']])

        obj_ws = sp.worksheet('Объекты')
        obj_rows = obj_ws.get_all_values()[1:]
        table = []
        for r in obj_rows:
            if not r or not r[0]:
                continue
            pname = r[0]
            p_ops = [o for o in ops if o['project'] == pname]
            plan_rev = _num(r[4]) if len(r) > 4 else 0
            plan_exp = _num(r[5]) if len(r) > 5 else 0
            plan_prf = plan_rev - plan_exp
            plan_mrg = plan_prf / plan_rev if plan_rev > 0 else 0
            f_inc = sum(o['amount'] for o in p_ops if o['type'] == 'Приход')
            f_exp = sum(o['amount'] for o in p_ops if o['type'] == 'Расход')
            f_prf = f_inc - f_exp
            f_mrg = f_prf / f_inc if f_inc > 0 else 0
            table.append([pname, r[1] if len(r)>1 else '',
                plan_rev, plan_exp, plan_prf, round(plan_mrg*100,1),
                f_inc, f_exp, f_prf, round(f_mrg*100,1), f_prf - plan_prf, f_prf])
        if table:
            ws.update(f'A10:L{9+len(table)}', table)
    except Exception as e:
        logger.error(f"Dashboard refresh error: {e}")

# ── Cash Flow ─────────────────────────────────────────────────────────────────

def _refresh_cashflow(sp):
    try:
        ws = sp.worksheet('Cash Flow')
        ws.clear()
        ops = _get_operations_raw(sp)

        months = []
        now = datetime.now()
        for i in range(-3, 9):
            m = now.month + i
            y = now.year + (m - 1) // 12
            m = ((m - 1) % 12) + 1
            months.append((y, m))

        header = ['Показатель'] + [f"{m:02d}.{y}" for y, m in months]

        def month_sum(y, m, t):
            return sum(o['amount'] for o in ops
                if o['type'] == t and _parse_date(o['date']) and
                _parse_date(o['date']).year == y and _parse_date(o['date']).month == m)

        income_row  = ['Приход']
        expense_row = ['Расход']
        saldo_row   = ['Сальдо']
        cumul_row   = ['Накопительно']
        margin_row  = ['Маржа']
        cumul = 0
        for y, m in months:
            inc = month_sum(y, m, 'Приход')
            exp = month_sum(y, m, 'Расход')
            sal = inc - exp
            cumul += sal
            income_row.append(inc)
            expense_row.append(exp)
            saldo_row.append(sal)
            cumul_row.append(cumul)
            margin_row.append(sal)

        ws.update('A1', [header, income_row, expense_row, saldo_row, cumul_row, margin_row])
    except Exception as e:
        logger.error(f"Cash flow error: {e}")

# ── Public data functions ─────────────────────────────────────────────────────

def get_projects():
    return _get_projects_raw(get_sheets())

def get_project_summary(project):
    sp = get_sheets()
    ops = _get_operations_raw(sp)
    p_ops = [o for o in ops if o['project'] == project]
    _, row = _get_project_row_raw(sp, project)
    if row is None:
        return None
    plan_rev = _num(row[4]) if len(row)>4 else 0
    plan_exp = _num(row[5]) if len(row)>5 else 0
    plan_profit  = plan_rev - plan_exp
    plan_margin  = plan_profit / plan_rev if plan_rev > 0 else 0
    f_inc  = sum(o['amount'] for o in p_ops if o['type'] == 'Приход')
    f_exp  = sum(o['amount'] for o in p_ops if o['type'] == 'Расход')
    f_profit  = f_inc - f_exp
    f_margin  = f_profit / f_inc if f_inc > 0 else 0
    return {
        'name': project, 'status': row[1] if len(row)>1 else '',
        'plan_revenue': plan_rev, 'plan_expense': plan_exp,
        'plan_profit': plan_profit, 'plan_margin': plan_margin,
        'fact_income': f_inc, 'fact_expense': f_exp,
        'fact_profit': f_profit, 'fact_margin': f_margin,
        'margin': f_profit,
        'dev_profit': f_profit - plan_profit,
        'recent_ops': sorted(p_ops, key=lambda x: _parse_date(x['date']) or datetime.min, reverse=True),
    }

def get_all_summary():
    sp = get_sheets()
    ops = _get_operations_raw(sp)
    obj_rows = sp.worksheet('Объекты').get_all_values()[1:]
    f_inc  = sum(o['amount'] for o in ops if o['type'] == 'Приход')
    f_exp  = sum(o['amount'] for o in ops if o['type'] == 'Расход')
    f_profit  = f_inc - f_exp
    f_margin  = f_profit / f_inc if f_inc > 0 else 0
    plan_rev = sum(_num(r[4]) for r in obj_rows if r and r[0] and len(r)>4)
    plan_exp = sum(_num(r[5]) for r in obj_rows if r and r[0] and len(r)>5)
    plan_profit  = plan_rev - plan_exp
    plan_margin  = plan_profit / plan_rev if plan_rev > 0 else 0
    return {'fact_income':f_inc, 'fact_expense':f_exp,
            'fact_profit':f_profit, 'fact_margin':f_margin,
            'plan_revenue':plan_rev, 'plan_expense':plan_exp,
            'plan_profit':plan_profit, 'plan_margin':plan_margin,
            'margin': f_profit,
            'projects_count':len([r for r in obj_rows if r and r[0]])}

def get_margin_report():
    sp = get_sheets()
    ops = _get_operations_raw(sp)
    obj_rows = sp.worksheet('Объекты').get_all_values()[1:]
    projects = []
    total_margin = 0
    for r in obj_rows:
        if not r or not r[0]:
            continue
        pname = r[0]
        p_ops = [o for o in ops if o['project'] == pname]
        f_inc = sum(o['amount'] for o in p_ops if o['type'] == 'Приход')
        f_exp = sum(o['amount'] for o in p_ops if o['type'] == 'Расход')
        margin = f_inc - f_exp
        margin_pct = margin / f_inc if f_inc > 0 else 0
        total_margin += margin
        projects.append({'name': pname, 'income': f_inc, 'expense': f_exp,
                         'margin': margin, 'margin_pct': margin_pct})
    projects.sort(key=lambda x: x['margin'], reverse=True)
    return {'total_margin': total_margin, 'projects': projects}

def add_operation(project, op_type, category, amount, date, pay_status, contractor='', comment=''):
    sp = get_sheets()
    ws = sp.worksheet('Операции')
    next_id = _next_op_id(sp)
    ws.append_row([next_id, date, project, op_type, _norm_cat(category),
                   amount, contractor, pay_status, comment])
    sb_mirror_operation(next_id, project, op_type, category, amount, date, pay_status, contractor, comment)
    _update_object_sheet(sp, project)
    _refresh_dashboard(sp)
    _refresh_cashflow(sp)
    warning = None
    _, proj_row = _get_project_row_raw(sp, project)
    if proj_row and op_type == 'Расход':
        plan_exp = _num(proj_row[5]) if len(proj_row)>5 else 0
        if plan_exp > 0:
            ops = _get_operations_raw(sp)
            total_exp = sum(o['amount'] for o in ops if o['project']==project and o['type']=='Расход')
            pct = total_exp / plan_exp
            if pct >= 0.8:
                warning = f"⚠️ Расходы {project}: {round(pct*100):.0f}% от плана ({_fmt(total_exp)} / {_fmt(plan_exp)} ₽)"
    return next_id, warning

def create_project(name, revenue, plan_expense):
    sp = get_sheets()
    plan_profit = revenue - plan_expense
    plan_margin = plan_profit / revenue if revenue > 0 else 0
    sp.worksheet('Объекты').append_row([
        name, 'Активный', datetime.now().strftime('%d.%m.%Y'), '',
        revenue, plan_expense, plan_profit, round(plan_margin*100, 1), 0, 0])
    sb_mirror_project(name, revenue)
    _create_object_sheet(sp, name)
    _refresh_dashboard(sp)

def delete_operation(op_id):
    sp = get_sheets()
    ws = sp.worksheet('Операции')
    rows = ws.get_all_values()
    for i, row in enumerate(rows[1:], 2):
        if row and str(row[0]) == str(op_id):
            project = row[2] if len(row) > 2 else None
            ws.delete_rows(i)
            sb_mirror_delete(op_id)
            if project:
                _update_object_sheet(sp, project)
                _refresh_dashboard(sp)
                _refresh_cashflow(sp)
            return True
    return False

def get_recent_ops(project=None, limit=10):
    ops = _get_operations_raw(get_sheets())
    if project:
        ops = [o for o in ops if o['project'] == project]
    return sorted(ops, key=lambda x: x['date'], reverse=True)[:limit]

# ── Import from Excel ─────────────────────────────────────────────────────────

def parse_xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    result = {'projects': [], 'total_ops': 0}

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        if ws.max_row is None or ws.max_row <= 1:
            continue

        ops = []
        contract_value = 0

        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=12, values_only=False):
            a = row[0].value if len(row) > 0 else None
            b = row[1].value if len(row) > 1 else None
            c = row[2].value if len(row) > 2 else None
            d = row[3].value if len(row) > 3 else None
            e = row[4].value if len(row) > 4 else None
            f = row[5].value if len(row) > 5 else None

            if isinstance(a, datetime) and isinstance(b, (int, float)) and b != 0:
                op_type = 'Приход' if b > 0 else 'Расход'
                category = _norm_cat(str(c)) if c else 'Прочее'
                comment = str(d).strip() if d else ''
                contractor = str(e).strip() if e else ''
                if f and isinstance(f, str):
                    comment = f"{comment} ({f.strip()})" if comment else f.strip()
                ops.append({
                    'date': a.strftime('%d.%m.%Y'),
                    'amount': abs(b),
                    'type': op_type,
                    'category': category,
                    'comment': comment,
                    'contractor': contractor,
                })

            for cell in row:
                if cell.value and isinstance(cell.value, str):
                    lower = cell.value.lower().strip()
                    if any(k in lower for k in ('итого', 'договор')):
                        nc = ws.cell(cell.row, cell.column + 1)
                        if nc.value and isinstance(nc.value, (int, float)) and nc.value > 0:
                            contract_value = max(contract_value, nc.value)

        if ops:
            total_income = sum(o['amount'] for o in ops if o['type'] == 'Приход')
            total_expense = sum(o['amount'] for o in ops if o['type'] == 'Расход')
            if contract_value == 0:
                contract_value = total_income

            result['projects'].append({
                'name': sheet_name,
                'ops': ops,
                'contract_value': contract_value,
                'plan_expense': total_expense,
            })
            result['total_ops'] += len(ops)

    return result

def execute_import(data):
    sp = get_sheets()

    ws_ops = sp.worksheet('Операции')
    ws_ops.clear()
    ws_ops.update('A1:I1', [['ID','Дата','Объект','Тип','Категория',
        'Сумма','Контрагент','Статус оплаты','Комментарий']])

    ws_obj = sp.worksheet('Объекты')
    ws_obj.clear()
    ws_obj.update('A1:J1', [['Название','Статус','Дата начала','Адрес',
        'Сумма договора','План расход','План прибыль','План маржа %',
        'Факт доход','Факт расход']])

    titles = _ws_titles(sp)
    base_sheets = {'Объекты', 'Операции', 'Дашборд', 'Cash Flow'}
    for title in titles:
        if title not in base_sheets:
            try:
                sp.del_worksheet(sp.worksheet(title))
            except:
                pass

    op_id = 1
    all_ops = []

    for proj in data['projects']:
        plan_profit = proj['contract_value'] - proj['plan_expense']
        plan_margin = plan_profit / proj['contract_value'] if proj['contract_value'] > 0 else 0
        f_inc = sum(o['amount'] for o in proj['ops'] if o['type'] == 'Приход')
        f_exp = sum(o['amount'] for o in proj['ops'] if o['type'] == 'Расход')
        ws_obj.append_row([
            proj['name'], 'Активный', datetime.now().strftime('%d.%m.%Y'), '',
            proj['contract_value'], proj['plan_expense'],
            plan_profit, round(plan_margin * 100, 1), f_inc, f_exp
        ])

        for op in sorted(proj['ops'], key=lambda x: x['date']):
            all_ops.append([
                op_id, op['date'], proj['name'], op['type'], op['category'],
                op['amount'], op['contractor'], 'Оплачено', op['comment']
            ])
            op_id += 1

    if all_ops:
        ws_ops.update(f'A2:I{1+len(all_ops)}', all_ops)

    for proj in data['projects']:
        _create_object_sheet(sp, proj['name'])

    _refresh_dashboard(sp)
    _refresh_cashflow(sp)

    return op_id - 1

# ── Keyboards ─────────────────────────────────────────────────────────────────

def main_kb(is_admin=True):
    kb = []
    if is_admin:
        kb.append([InlineKeyboardButton("💸 Быстрый расход", callback_data='quick_expense')])
        kb.append([InlineKeyboardButton("➕ Новая операция", callback_data='op_new')])
    kb.append([InlineKeyboardButton("💰 Мой доход", callback_data='my_income'),
               InlineKeyboardButton("📊 Отчёт", callback_data='report_all')])
    kb.append([InlineKeyboardButton("🏢 По объекту", callback_data='report_object')])
    if is_admin:
        kb.append([InlineKeyboardButton("📁 Новый объект", callback_data='create_project'),
                   InlineKeyboardButton("✏️ Операции", callback_data='edit_ops')])
        kb.append([InlineKeyboardButton("📥 CSV", callback_data='export'),
                   InlineKeyboardButton("📤 Импорт xlsx", callback_data='import_start')])
    sheets_url = _sheets_url()
    if sheets_url:
        kb.append([InlineKeyboardButton("📋 Google Таблица", url=sheets_url)])
    return InlineKeyboardMarkup(kb)

def back_kb():
    return InlineKeyboardMarkup([[InlineKeyboardButton("◀ Назад", callback_data='back_to_menu')]])

OPS_PAGE_SIZE = 8

def _sheets_url():
    sid = os.environ.get('GOOGLE_SHEETS_ID', '')
    return f"https://docs.google.com/spreadsheets/d/{sid}" if sid else None

async def _show_object_report(q, context, s, page=0):
    sign = "+" if s['margin'] >= 0 else ""
    text = (f"📊 <b>{s['name']}</b> ({s['status']})\n\n"
            f"💰 <b>Маржа (мой доход): {sign}{_fmt(s['margin'])} ₽ ({_pct(s['fact_margin'])})</b>\n\n"
            f"<b>ПЛАН:</b>\n"
            f"💰 Договор: {_fmt(s['plan_revenue'])} ₽\n"
            f"💸 Расход: {_fmt(s['plan_expense'])} ₽\n"
            f"📈 Прибыль: {_fmt(s['plan_profit'])} ₽\n"
            f"📊 Маржа: {_pct(s['plan_margin'])}\n\n"
            f"<b>ФАКТ:</b>\n"
            f"💰 Доход: {_fmt(s['fact_income'])} ₽\n"
            f"💸 Расход: {_fmt(s['fact_expense'])} ₽\n"
            f"📈 Прибыль: {_fmt(s['fact_profit'])} ₽\n"
            f"📊 Маржа: {_pct(s['fact_margin'])}\n\n"
            f"📉 Откл.: {_fmt(s['dev_profit'])} ₽\n")

    ops = s['recent_ops']
    total_pages = max(1, (len(ops) + OPS_PAGE_SIZE - 1) // OPS_PAGE_SIZE)
    page = max(0, min(page, total_pages - 1))
    chunk = ops[page * OPS_PAGE_SIZE : (page + 1) * OPS_PAGE_SIZE]

    if ops:
        text += f"\n<b>Операции (стр. {page + 1}/{total_pages}):</b>\n"
        for o in chunk:
            sign_op = "+" if o['type'] == 'Приход' else "-"
            text += f"  {o['date']}  {sign_op}{_fmt(o['amount'])} ₽  {o['category']}\n"

    kb = []
    if total_pages > 1:
        nav = []
        if page > 0:
            nav.append(InlineKeyboardButton("◀ Пред", callback_data=f"rptpg_{page - 1}"))
        nav.append(InlineKeyboardButton(f"{page + 1}/{total_pages}", callback_data="noop"))
        if page < total_pages - 1:
            nav.append(InlineKeyboardButton("След ▶", callback_data=f"rptpg_{page + 1}"))
        kb.append(nav)
    kb.append([InlineKeyboardButton("◀ Назад", callback_data='back_to_menu')])

    await q.edit_message_text(text, reply_markup=InlineKeyboardMarkup(kb), parse_mode=ParseMode.HTML)

# ── Handlers ──────────────────────────────────────────────────────────────────

@any_user
async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.clear()
    admin = _is_admin(update)
    await update.message.reply_text(
        "🌱 ZenScape — финансовый учёт\n\nЧто хочешь сделать?",
        reply_markup=main_kb(admin))
    return S.MENU

@any_user
async def menu_cb(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query
    await q.answer()
    d = q.data
    admin = _is_admin(update)

    # ── Back to menu ──
    if d == 'back_to_menu':
        context.user_data.clear()
        await q.edit_message_text("Что хочешь сделать?", reply_markup=main_kb(admin))
        return S.MENU

    if d == 'noop':
        return S.MENU

    # ── Quick expense ──
    if d == 'quick_expense' and admin:
        projects = get_projects()
        if not projects:
            await q.edit_message_text("Нет объектов. Создай объект сначала.", reply_markup=back_kb())
            return S.MENU
        kb = [[InlineKeyboardButton(p, callback_data=f'qp_{p}')] for p in projects]
        kb.append([InlineKeyboardButton("◀ Назад", callback_data='back_to_menu')])
        await q.edit_message_text("💸 Быстрый расход\n\nВыбери объект:", reply_markup=InlineKeyboardMarkup(kb))
        return S.QUICK_PROJECT

    if d.startswith('qp_'):
        context.user_data['project'] = d[3:]
        await q.edit_message_text(f"💸 {context.user_data['project']}\n\nСумма расхода (₽):")
        return S.QUICK_AMOUNT

    # ── New operation ──
    if d == 'op_new' and admin:
        projects = get_projects()
        if not projects:
            await q.edit_message_text("Нет объектов. Создай объект сначала.", reply_markup=back_kb())
            return S.MENU
        kb = [[InlineKeyboardButton(p, callback_data=f'proj_{p}')] for p in projects]
        kb.append([InlineKeyboardButton("◀ Назад", callback_data='back_to_menu')])
        await q.edit_message_text("Выбери объект:", reply_markup=InlineKeyboardMarkup(kb))
        return S.SELECT_PROJECT

    if d.startswith('proj_'):
        context.user_data['project'] = d[5:]
        await q.edit_message_text(
            f"📍 {context.user_data['project']}\n\nТип операции:",
            reply_markup=InlineKeyboardMarkup([
                [InlineKeyboardButton("💰 Приход", callback_data='type_Приход')],
                [InlineKeyboardButton("💸 Расход", callback_data='type_Расход')],
                [InlineKeyboardButton("◀ Назад", callback_data='back_to_menu')],
            ]))
        return S.OP_TYPE

    if d.startswith('type_'):
        context.user_data['op_type'] = d[5:]
        cats = INCOME_CATS if d[5:] == 'Приход' else EXPENSE_CATS
        kb = [[InlineKeyboardButton(c, callback_data=f'cat_{c}')] for c in cats]
        kb.append([InlineKeyboardButton("◀ Назад", callback_data='back_to_menu')])
        await q.edit_message_text(f"Категория ({d[5:]}):", reply_markup=InlineKeyboardMarkup(kb))
        return S.CATEGORY

    if d.startswith('cat_'):
        context.user_data['category'] = d[4:]
        await q.edit_message_text("Сумма (₽):")
        return S.AMOUNT

    if d in ('date_today','date_yesterday','date_custom'):
        if d == 'date_today':
            context.user_data['date'] = datetime.now().strftime('%d.%m.%Y')
        elif d == 'date_yesterday':
            context.user_data['date'] = (datetime.now()-timedelta(days=1)).strftime('%d.%m.%Y')
        else:
            await q.edit_message_text("Напиши дату в формате ДД.ММ.ГГГГ:")
            return S.DATE
        return await _show_pay_status(q, context)

    if d.startswith('pay_'):
        context.user_data['pay_status'] = d[4:]
        await q.edit_message_text("Контрагент:",
            reply_markup=InlineKeyboardMarkup([[InlineKeyboardButton("⏭ Пропустить", callback_data='skip_contractor')]]))
        return S.CONTRACTOR

    if d == 'skip_contractor':
        context.user_data['contractor'] = ''
        await q.edit_message_text("Комментарий:",
            reply_markup=InlineKeyboardMarkup([[InlineKeyboardButton("⏭ Пропустить", callback_data='skip_comment')]]))
        return S.COMMENT

    if d == 'skip_comment':
        context.user_data['comment'] = ''
        return await _save_op_q(q, context)

    # ── My income (margin) ──
    if d == 'my_income':
        report = get_margin_report()
        text = f"💰 <b>Мой доход</b>\n\n"
        text += f"На руках: <b>{_fmt(report['total_margin'])} ₽</b>\n\n"
        for p in report['projects']:
            sign = "+" if p['margin'] >= 0 else ""
            emoji = "🟢" if p['margin'] > 0 else "🔴"
            text += f"{emoji} <b>{p['name']}</b>\n"
            text += f"   {sign}{_fmt(p['margin'])} ₽"
            if p['margin_pct'] != 0:
                text += f" (маржа {_pct(p['margin_pct'])})"
            text += "\n"
            text += f"   ↳ доход {_fmt(p['income'])} / расход {_fmt(p['expense'])} ₽\n\n"
        await q.edit_message_text(text, reply_markup=back_kb(), parse_mode=ParseMode.HTML)
        return S.MENU

    # ── Reports ──
    if d == 'report_all':
        s = get_all_summary()
        text = (f"📊 <b>Общий отчёт</b>\n\nОбъектов: {s['projects_count']}\n\n"
                f"<b>ФАКТ:</b>\n"
                f"💰 Доход: {_fmt(s['fact_income'])} ₽\n"
                f"💸 Расход: {_fmt(s['fact_expense'])} ₽\n"
                f"📈 Прибыль: {_fmt(s['fact_profit'])} ₽\n"
                f"📊 Маржа: {_pct(s['fact_margin'])}\n"
                f"💰 Мой доход: <b>{_fmt(s['margin'])} ₽</b>\n\n"
                f"<b>ПЛАН:</b>\n"
                f"💰 Выручка: {_fmt(s['plan_revenue'])} ₽\n"
                f"💸 Расход: {_fmt(s['plan_expense'])} ₽\n"
                f"📈 Прибыль: {_fmt(s['plan_profit'])} ₽\n"
                f"📊 Маржа: {_pct(s['plan_margin'])}")
        await q.edit_message_text(text, reply_markup=back_kb(), parse_mode=ParseMode.HTML)
        return S.MENU

    if d == 'report_object':
        projects = get_projects()
        kb = [[InlineKeyboardButton(p, callback_data=f'rpt_{p}')] for p in projects]
        kb.append([InlineKeyboardButton("◀ Назад", callback_data='back_to_menu')])
        await q.edit_message_text("Выбери объект:", reply_markup=InlineKeyboardMarkup(kb))
        return S.MENU

    if d.startswith('rpt_'):
        project = d[4:]
        s = get_project_summary(project)
        if not s:
            await q.edit_message_text("Объект не найден.", reply_markup=back_kb())
            return S.MENU
        context.user_data['report_project'] = project
        context.user_data['report_ops'] = s['recent_ops']
        await _show_object_report(q, context, s, page=0)
        return S.MENU

    if d.startswith('rptpg_'):
        page = int(d[6:])
        project = context.user_data.get('report_project')
        s = get_project_summary(project) if project else None
        if not s:
            await q.edit_message_text("Объект не найден.", reply_markup=back_kb())
            return S.MENU
        context.user_data['report_ops'] = s['recent_ops']
        await _show_object_report(q, context, s, page=page)
        return S.MENU

    # ── Create project ──
    if d == 'create_project' and admin:
        await q.edit_message_text("Название нового объекта:")
        return S.CP_NAME

    # ── Edit operations ──
    if d == 'edit_ops' and admin:
        ops = get_recent_ops(limit=10)
        if not ops:
            await q.edit_message_text("Операций нет.", reply_markup=back_kb())
            return S.MENU
        kb = []
        for o in ops:
            sign = "+" if o['type'] == 'Приход' else "-"
            label = f"{o['date']} {o['project'][:10]} {sign}{_fmt(o['amount'])}₽"
            kb.append([InlineKeyboardButton(f"🗑 {label}", callback_data=f'del_{o["id"]}')])
        kb.append([InlineKeyboardButton("◀ Назад", callback_data='back_to_menu')])
        await q.edit_message_text("Нажми 🗑 для удаления:", reply_markup=InlineKeyboardMarkup(kb))
        return S.EDIT_LIST

    if d.startswith('del_') and admin:
        if delete_operation(d[4:]):
            await q.answer("✅ Удалено", show_alert=True)
        else:
            await q.answer("❌ Ошибка", show_alert=True)
        ops = get_recent_ops(limit=10)
        if not ops:
            await q.edit_message_text("Операций больше нет.", reply_markup=back_kb())
            return S.MENU
        kb = []
        for o in ops:
            sign = "+" if o['type'] == 'Приход' else "-"
            label = f"{o['date']} {o['project'][:10]} {sign}{_fmt(o['amount'])}₽"
            kb.append([InlineKeyboardButton(f"🗑 {label}", callback_data=f'del_{o["id"]}')])
        kb.append([InlineKeyboardButton("◀ Назад", callback_data='back_to_menu')])
        await q.edit_message_text("Нажми 🗑 для удаления:", reply_markup=InlineKeyboardMarkup(kb))
        return S.EDIT_LIST

    # ── Export ──
    if d == 'export' and admin:
        projects = get_projects()
        kb = [[InlineKeyboardButton("📦 Все объекты", callback_data='exp_all')]]
        kb += [[InlineKeyboardButton(p, callback_data=f'expp_{p}')] for p in projects]
        kb.append([InlineKeyboardButton("◀ Назад", callback_data='back_to_menu')])
        await q.edit_message_text("Выгрузить:", reply_markup=InlineKeyboardMarkup(kb))
        return S.MENU

    if d in ('exp_all',) or d.startswith('expp_'):
        project = None if d == 'exp_all' else d[5:]
        ops = _get_operations_raw(get_sheets())
        if project:
            ops = [o for o in ops if o['project'] == project]
        out = io.StringIO()
        w = csv.writer(out)
        w.writerow(['ID','Дата','Объект','Тип','Категория','Сумма','Контрагент','Статус оплаты','Комментарий'])
        for o in ops:
            w.writerow([o['id'],o['date'],o['project'],o['type'],o['category'],
                        o['amount'],o['contractor'],o['pay_status'],o['comment']])
        bio = io.BytesIO(out.getvalue().encode('utf-8-sig'))
        fname = f"zenscape_{project or 'all'}.csv"
        await q.message.reply_document(document=bio, filename=fname, caption="✅ Готово")
        await q.edit_message_text("✅ Файл отправлен.", reply_markup=back_kb())
        return S.MENU

    # ── Import start ──
    if d == 'import_start' and admin:
        await q.edit_message_text(
            "📤 <b>Импорт из Excel</b>\n\n"
            "Отправь файл .xlsx с данными.\n\n"
            "Каждый лист = один объект.\n"
            "Столбцы: Дата | Сумма | Статья | Описание | Контрагент\n\n"
            "⚠️ Текущие данные будут заменены!",
            parse_mode=ParseMode.HTML, reply_markup=back_kb())
        return S.IMPORT_FILE

    if d == 'import_yes':
        data = context.user_data.get('import_data')
        if not data:
            await q.edit_message_text("❌ Нет данных для импорта.", reply_markup=back_kb())
            return S.MENU
        try:
            total = execute_import(data)
            await q.edit_message_text(
                f"✅ Импорт завершён!\n\n"
                f"📁 Объектов: {len(data['projects'])}\n"
                f"📝 Операций: {total}\n\n"
                f"Дашборд и листы объектов обновлены.",
                reply_markup=main_kb(admin))
            context.user_data.clear()
        except Exception as e:
            logger.error(f"Import error: {e}")
            await q.edit_message_text(f"❌ Ошибка импорта: {e}", reply_markup=back_kb())
        return S.MENU

    if d == 'import_no':
        context.user_data.clear()
        await q.edit_message_text("Импорт отменён.", reply_markup=main_kb(admin))
        return S.MENU

    return S.MENU

async def _show_pay_status(q, context):
    await q.edit_message_text(
        f"📅 {context.user_data['date']}\n\nСтатус оплаты:",
        reply_markup=InlineKeyboardMarkup([
            [InlineKeyboardButton("✅ Оплачено", callback_data='pay_Оплачено')],
            [InlineKeyboardButton("⏳ Ожидает", callback_data='pay_Ожидает')],
            [InlineKeyboardButton("🔸 Частично", callback_data='pay_Частично')],
        ]))
    return S.PAY_STATUS

async def _save_op_q(q, context):
    op_id, warning = add_operation(
        context.user_data['project'], context.user_data['op_type'],
        context.user_data['category'], context.user_data['amount'],
        context.user_data.get('date', datetime.now().strftime('%d.%m.%Y')),
        context.user_data.get('pay_status', 'Оплачено'),
        context.user_data.get('contractor', ''), context.user_data.get('comment', ''))
    text = (f"✅ Сохранено!\n\n"
            f"📍 {context.user_data['project']}\n"
            f"{'💰' if context.user_data['op_type']=='Приход' else '💸'} "
            f"{context.user_data['op_type']} — {_fmt(context.user_data['amount'])} ₽\n"
            f"🏷 {context.user_data['category']}\n"
            f"📅 {context.user_data.get('date','')}\n"
            f"💳 {context.user_data.get('pay_status','')}")
    if warning:
        text += f"\n\n{warning}"
    admin = q.from_user.id in ADMIN_USERS
    context.user_data.clear()
    await q.edit_message_text(text, reply_markup=main_kb(admin))
    return S.MENU

async def _save_op_msg(update, context):
    op_id, warning = add_operation(
        context.user_data['project'], context.user_data['op_type'],
        context.user_data['category'], context.user_data['amount'],
        context.user_data.get('date', datetime.now().strftime('%d.%m.%Y')),
        context.user_data.get('pay_status', 'Оплачено'),
        context.user_data.get('contractor', ''), context.user_data.get('comment', ''))
    text = (f"✅ Сохранено!\n\n"
            f"📍 {context.user_data['project']}\n"
            f"{'💰' if context.user_data['op_type']=='Приход' else '💸'} "
            f"{context.user_data['op_type']} — {_fmt(context.user_data['amount'])} ₽\n"
            f"🏷 {context.user_data['category']}\n"
            f"📅 {context.user_data.get('date','')}\n"
            f"💳 {context.user_data.get('pay_status','')}")
    if warning:
        text += f"\n\n{warning}"
    admin = update.effective_user.id in ADMIN_USERS
    context.user_data.clear()
    await update.message.reply_text(text, reply_markup=main_kb(admin))
    return S.MENU

# ── Text input handlers ───────────────────────────────────────────────────────

@admin_only
async def quick_amount_handler(update, context):
    try:
        amount = float(update.message.text.strip().replace(',','.').replace(' ',''))
        op_id, warning = add_operation(
            context.user_data['project'], 'Расход', 'Прочее', amount,
            datetime.now().strftime('%d.%m.%Y'), 'Оплачено')
        text = f"✅ Расход записан!\n\n📍 {context.user_data['project']}\n💸 {_fmt(amount)} ₽"
        if warning:
            text += f"\n\n{warning}"
        context.user_data.clear()
        await update.message.reply_text(text, reply_markup=main_kb(True))
        return S.MENU
    except ValueError:
        await update.message.reply_text("❌ Введи число, например: 15000")
        return S.QUICK_AMOUNT

@admin_only
async def amount_handler(update, context):
    try:
        context.user_data['amount'] = float(update.message.text.strip().replace(',','.').replace(' ',''))
        today = datetime.now()
        yesterday = today - timedelta(days=1)
        await update.message.reply_text("Дата операции:", reply_markup=InlineKeyboardMarkup([
            [InlineKeyboardButton(f"📅 Сегодня ({today.strftime('%d.%m')})", callback_data='date_today')],
            [InlineKeyboardButton(f"📅 Вчера ({yesterday.strftime('%d.%m')})", callback_data='date_yesterday')],
            [InlineKeyboardButton("✏️ Своя дата", callback_data='date_custom')],
        ]))
        return S.DATE
    except ValueError:
        await update.message.reply_text("❌ Введи число, например: 15000")
        return S.AMOUNT

@admin_only
async def date_text_handler(update, context):
    try:
        datetime.strptime(update.message.text.strip(), '%d.%m.%Y')
        context.user_data['date'] = update.message.text.strip()
        await update.message.reply_text(
            f"📅 {context.user_data['date']}\n\nСтатус оплаты:",
            reply_markup=InlineKeyboardMarkup([
                [InlineKeyboardButton("✅ Оплачено", callback_data='pay_Оплачено')],
                [InlineKeyboardButton("⏳ Ожидает", callback_data='pay_Ожидает')],
                [InlineKeyboardButton("🔸 Частично", callback_data='pay_Частично')],
            ]))
        return S.PAY_STATUS
    except ValueError:
        await update.message.reply_text("❌ Формат: ДД.ММ.ГГГГ")
        return S.DATE

@admin_only
async def contractor_handler(update, context):
    context.user_data['contractor'] = update.message.text.strip()
    await update.message.reply_text("Комментарий:",
        reply_markup=InlineKeyboardMarkup([[InlineKeyboardButton("⏭ Пропустить", callback_data='skip_comment')]]))
    return S.COMMENT

@admin_only
async def comment_handler(update, context):
    context.user_data['comment'] = update.message.text.strip()
    return await _save_op_msg(update, context)

@admin_only
async def cp_name_handler(update, context):
    context.user_data['cp_name'] = update.message.text.strip()
    await update.message.reply_text(
        f"Объект: <b>{context.user_data['cp_name']}</b>\n\nСумма договора (₽):",
        parse_mode=ParseMode.HTML)
    return S.CP_REVENUE

@admin_only
async def cp_revenue_handler(update, context):
    try:
        context.user_data['cp_revenue'] = float(update.message.text.strip().replace(',','.').replace(' ',''))
        await update.message.reply_text(
            f"Выручка: {_fmt(context.user_data['cp_revenue'])} ₽\n\nПлановые расходы (₽):\n<i>0 если неизвестно</i>",
            parse_mode=ParseMode.HTML)
        return S.CP_EXPENSE
    except ValueError:
        await update.message.reply_text("❌ Введи число")
        return S.CP_REVENUE

@admin_only
async def cp_expense_handler(update, context):
    try:
        expense = float(update.message.text.strip().replace(',','.').replace(' ',''))
        name = context.user_data['cp_name']
        revenue = context.user_data['cp_revenue']
        profit = revenue - expense
        margin = profit / revenue if revenue > 0 else 0
        create_project(name, revenue, expense)
        await update.message.reply_text(
            f"✅ Объект создан!\n\n"
            f"📍 {name}\n"
            f"💰 Договор: {_fmt(revenue)} ₽\n"
            f"💸 Расход: {_fmt(expense)} ₽\n"
            f"📈 Прибыль: {_fmt(profit)} ₽\n"
            f"📊 Маржа: {_pct(margin)}\n\n"
            f"Вкладка в таблице создана автоматически.",
            reply_markup=main_kb(True))
        context.user_data.clear()
        return S.MENU
    except ValueError:
        await update.message.reply_text("❌ Введи число")
        return S.CP_EXPENSE

@admin_only
async def import_file_handler(update, context):
    doc = update.message.document
    if not doc or not doc.file_name.endswith('.xlsx'):
        await update.message.reply_text("❌ Отправь файл .xlsx", reply_markup=back_kb())
        return S.IMPORT_FILE

    await update.message.reply_text("⏳ Обрабатываю файл...")

    try:
        file = await doc.get_file()
        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
            await file.download_to_drive(tmp.name)
            data = parse_xlsx(tmp.name)
            os.unlink(tmp.name)
    except Exception as e:
        logger.error(f"Import parse error: {e}")
        await update.message.reply_text(f"❌ Ошибка чтения файла: {e}", reply_markup=back_kb())
        return S.MENU

    if not data['projects']:
        await update.message.reply_text("❌ Не найдено данных для импорта.", reply_markup=back_kb())
        return S.MENU

    context.user_data['import_data'] = data

    text = f"📤 <b>Найдено для импорта:</b>\n\n"
    for p in data['projects']:
        inc = sum(o['amount'] for o in p['ops'] if o['type'] == 'Приход')
        exp = sum(o['amount'] for o in p['ops'] if o['type'] == 'Расход')
        text += f"📍 <b>{p['name']}</b>\n"
        text += f"   Операций: {len(p['ops'])}\n"
        text += f"   Доход: {_fmt(inc)} ₽ | Расход: {_fmt(exp)} ₽\n"
        if p['contract_value']:
            text += f"   Договор: {_fmt(p['contract_value'])} ₽\n"
        text += "\n"
    text += f"📝 Всего операций: {data['total_ops']}\n\n"
    text += "⚠️ <b>Текущие данные будут УДАЛЕНЫ.</b>\nПродолжить?"

    await update.message.reply_text(text, parse_mode=ParseMode.HTML,
        reply_markup=InlineKeyboardMarkup([
            [InlineKeyboardButton("✅ Да, импортировать", callback_data='import_yes')],
            [InlineKeyboardButton("❌ Отмена", callback_data='import_no')],
        ]))
    return S.IMPORT_CONFIRM

@any_user
async def cancel(update, context):
    context.user_data.clear()
    admin = _is_admin(update)
    await update.message.reply_text("Отменено.", reply_markup=main_kb(admin))
    return S.MENU

# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    app = Application.builder().token(os.environ['TELEGRAM_BOT_TOKEN']).build()
    conv = ConversationHandler(
        entry_points=[CommandHandler('start', start)],
        states={
            S.MENU:           [CallbackQueryHandler(menu_cb)],
            S.SELECT_PROJECT: [CallbackQueryHandler(menu_cb)],
            S.OP_TYPE:        [CallbackQueryHandler(menu_cb)],
            S.CATEGORY:       [CallbackQueryHandler(menu_cb)],
            S.AMOUNT:         [MessageHandler(filters.TEXT & ~filters.COMMAND, amount_handler)],
            S.DATE:           [CallbackQueryHandler(menu_cb, pattern='^date_'),
                               MessageHandler(filters.TEXT & ~filters.COMMAND, date_text_handler)],
            S.PAY_STATUS:     [CallbackQueryHandler(menu_cb, pattern='^pay_')],
            S.CONTRACTOR:     [CallbackQueryHandler(menu_cb, pattern='^skip_contractor$'),
                               MessageHandler(filters.TEXT & ~filters.COMMAND, contractor_handler)],
            S.COMMENT:        [CallbackQueryHandler(menu_cb, pattern='^skip_comment$'),
                               MessageHandler(filters.TEXT & ~filters.COMMAND, comment_handler)],
            S.CP_NAME:        [MessageHandler(filters.TEXT & ~filters.COMMAND, cp_name_handler)],
            S.CP_REVENUE:     [MessageHandler(filters.TEXT & ~filters.COMMAND, cp_revenue_handler)],
            S.CP_EXPENSE:     [MessageHandler(filters.TEXT & ~filters.COMMAND, cp_expense_handler)],
            S.EDIT_LIST:      [CallbackQueryHandler(menu_cb)],
            S.IMPORT_FILE:    [MessageHandler(filters.Document.ALL, import_file_handler),
                               CallbackQueryHandler(menu_cb, pattern='^back_to_menu$')],
            S.IMPORT_CONFIRM: [CallbackQueryHandler(menu_cb)],
            S.QUICK_PROJECT:  [CallbackQueryHandler(menu_cb)],
            S.QUICK_AMOUNT:   [MessageHandler(filters.TEXT & ~filters.COMMAND, quick_amount_handler)],
        },
        fallbacks=[CommandHandler('cancel', cancel), CommandHandler('start', start)],
    )
    app.add_handler(conv)
    start_sync_thread()
    app.run_polling()

if __name__ == '__main__':
    main()
